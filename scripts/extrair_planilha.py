"""Extrai a planilha "Gerencial Financeiro" para JSON em data/import/.

Uso:
    python -I scripts/extrair_planilha.py "<caminho do .xlsx>"

Gera:
    data/import/plano_contas.json      contas (nome original -> conta padrão)
    data/import/lancamentos.json       linhas do razão, na ordem da planilha
    data/import/referencia.json        DRE e Balanço anuais calculados pelo Excel (gabarito)
"""
import datetime as dt
import json
import pathlib
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "data" / "import"


def corrigir_mojibake(s):
    """Históricos colados de extrato às vezes vêm com UTF-8 lido como Latin-1 ('TransferÃªncia')."""
    if "Ã" not in s and "Â" not in s:
        return s
    for codec in ("cp1252", "latin-1"):
        try:
            return s.encode(codec).decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
    return s


def texto(v):
    return None if v is None else corrigir_mojibake(str(v).strip())


def extrair_plano(ws):
    contas = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        original, padrao, categoria, atividade, natureza = row[:5]
        if not original:
            continue
        contas.append({
            "linha": i,
            "contaOriginal": texto(original),
            "contaPadrao": texto(padrao),
            "categoria": texto(categoria),
            "atividade": texto(atividade),
            "natureza": texto(natureza),
        })
    return contas


def extrair_lancamentos(ws):
    linhas, avisos = [], []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        data, original, padrao, categoria, atividade, historico, valor, fiscal, nominal, dc, ano, mes = row[:12]
        if data is None and valor is None:
            continue
        if not isinstance(data, dt.datetime):
            avisos.append(f"linha {i}: data não é data real ({data!r})")
            continue
        if not isinstance(valor, (int, float)):
            avisos.append(f"linha {i}: valor não numérico ({valor!r})")
            continue
        if ano is not None and int(ano) != data.year:
            avisos.append(f"linha {i}: coluna Ano={ano} difere do ano da data {data.date()}")
        linhas.append({
            "linha": i,
            "data": data.date().isoformat(),
            "contaOriginal": texto(original),
            "contaPadrao": texto(padrao),
            "categoria": texto(categoria),
            "atividade": texto(atividade),
            "historico": texto(historico),
            "valor": float(valor),
            "fiscal": texto(fiscal),
            "valorNominal": float(nominal) if isinstance(nominal, (int, float)) else None,
            "dc": texto(dc),
            "ano": int(ano) if ano is not None else data.year,
        })
    return linhas, avisos


def extrair_tabela_anual(ws):
    """Lê uma aba no formato: linha de anos + linhas 'rótulo | valor por ano'."""
    anos, tabela = None, {}
    for row in ws.iter_rows(values_only=True):
        if anos is None:
            if row and row[0] is None and isinstance(row[1], int):
                anos = [a for a in row[1:] if isinstance(a, int)]
            continue
        rotulo = texto(row[0])
        valores = row[1:1 + len(anos)]
        if rotulo and any(isinstance(v, (int, float)) for v in valores):
            tabela[rotulo] = {str(a): (v if isinstance(v, (int, float)) else None) for a, v in zip(anos, valores)}
    return tabela


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)

    plano = extrair_plano(wb["Plano de Contas"])
    lancamentos, avisos = extrair_lancamentos(wb["Lançamentos"])
    referencia = {
        "dre": extrair_tabela_anual(wb["DRE Anual"]),
        "balanco": extrair_tabela_anual(wb["Balanço Patrimonial Anual"]),
    }

    SAIDA.mkdir(parents=True, exist_ok=True)
    for nome, conteudo in [("plano_contas", plano), ("lancamentos", lancamentos), ("referencia", referencia)]:
        (SAIDA / f"{nome}.json").write_text(json.dumps(conteudo, ensure_ascii=False, indent=1), encoding="utf-8")

    originais = {c["contaOriginal"] for c in plano}
    nao_mapeadas = sorted({l["contaOriginal"] for l in lancamentos} - originais)
    sub_centavo = [l["linha"] for l in lancamentos if abs(round(l["valor"], 2) - l["valor"]) > 1e-9]

    print(f"Contas no plano: {len(plano)}")
    print(f"Lançamentos: {len(lancamentos)} ({lancamentos[0]['data']} a {lancamentos[-1]['data']})")
    print(f"Contas usadas sem mapeamento no plano: {nao_mapeadas or 'nenhuma'}")
    print(f"Linhas com fração de centavo: {len(sub_centavo)} {sub_centavo[:20]}")
    print(f"Avisos: {len(avisos)}")
    for a in avisos:
        print("  -", a)


if __name__ == "__main__":
    main()
